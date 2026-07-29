"""Build the formatted MATH6186 case-study report from its validated Markdown."""

from __future__ import annotations

import csv
import json
import math
import re
from collections import Counter
from pathlib import Path
from typing import Sequence

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips
from PIL import Image


PRESET = {
    "page_width_inches": 8.5,
    "page_height_inches": 11.0,
    "margin_inches": 1.0,
    "header_footer_inches": 0.492,
    "content_width_dxa": 9360,
    "body_font": "Calibri",
    "body_size_pt": 11,
    "body_after_pt": 8,
    "body_line_spacing": 1.333,
    "h1": (16, "2E74B5", 18, 10),
    "h2": (13, "2E74B5", 12, 6),
    "h3": (12, "1F4D78", 8, 4),
    "table_indent_dxa": 120,
    "table_cell_margins_dxa": (80, 80, 120, 120),
    "table_header_fill": "F4F6F9",
}

# Named narrative_proposal override for this formal long-form report.
EDITORIAL_COVER = {
    "title_size_pt": 28,
    "title_color": "1F4D78",
    "subtitle_size_pt": 14,
    "subtitle_color": "2E74B5",
    "top_whitespace_pt": 116,
}

CAPTION_STYLE = {
    "font": "Calibri",
    "size_pt": 9,
    "color": "595959",
    "before_pt": 4,
    "after_pt": 6,
    "line_spacing": 1.0,
}

ROOT = Path(__file__).resolve().parent.parent
DEFAULT_SOURCE = ROOT / "report" / "MATH6186_case_study_draft.md"
DEFAULT_OUTPUT = ROOT / "report" / "MATH6186_case_study_draft.docx"
REFERENCE_PATH = ROOT / "report" / "references.json"
SUMMARY_PATH = ROOT / "outputs" / "data" / "summary.csv"
PARAMETERS_PATH = ROOT / "outputs" / "data" / "parameters.csv"
FIGURE_DIR = ROOT / "outputs" / "figures"
RUNNING_HEADER = "MATH6186 | Worried-Well Consultation Capacity"

VALUE_PATTERN = re.compile(
    r"\[\[VALUE:([^|\]]+)\|([^|\]]+)\|([^|\]]+)\]\]"
)
FIGURE_PATTERN = re.compile(
    r"\[\[FIGURE:([^|\]]+)\|([^|\]]+)\|(fig[1-5])\]\]"
)
CITATION_PATTERN = re.compile(r"\[@([a-z0-9_]+)\]")
REFERENCE_PATTERN = re.compile(r"\[\[REFERENCE:([a-z0-9_]+)\]\]")
EQUATION_PATTERN = re.compile(r"\[\[EQUATION:([a-z_]+)\]\]")
DOUBLE_BRACKET_PATTERN = re.compile(r"\[\[[^\[\]\r\n]+\]\]")
PERMITTED_IDENTITY_TOKENS = {"[[STUDENT_NAME]]", "[[STUDENT_ID]]"}
EQUATIONS = {
    "compartment_system": [
        "dS/dt = -(beta_P + beta_WP)SI_P - beta_W SI_W + delta_P R_P + gamma_W I_W",
        "dI_P/dt = beta_P SI_P + alpha beta_P I_P I_W - gamma_P I_P",
        "dI_W/dt = -alpha beta_P I_P I_W + beta_W SI_W + beta_WP SI_P - gamma_W I_W",
        "dR_P/dt = gamma_P I_P - delta_P R_P",
    ],
    "expected_demand": ["m_t = N[p_P I_P(t) + p_W I_W(t)]"],
    "sample_average_cost": [
        "min_q (1/n) sum_j [c_u(D_j-q)^+ + c_o(q-D_j)^+]"
    ],
    "critical_fractile": ["tau = c_u/(c_u+c_o), q* = F_n^{-1}(tau)"],
}
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"
CANONICAL_SOURCE_TEXT = DEFAULT_SOURCE.read_text(encoding="utf-8")
CANONICAL_VALUE_MARKERS = frozenset(
    match.group(0) for match in VALUE_PATTERN.finditer(CANONICAL_SOURCE_TEXT)
)
CANONICAL_FIGURE_MARKERS = frozenset(
    match.group(0) for match in FIGURE_PATTERN.finditer(CANONICAL_SOURCE_TEXT)
)


def load_references(path: Path) -> dict[str, dict[str, object]]:
    """Load the canonical reference list keyed by its stable citation IDs."""

    with path.open(encoding="utf-8") as stream:
        payload = json.load(stream)
    if not isinstance(payload, list):
        raise ValueError("references.json must contain a JSON array.")
    references: dict[str, dict[str, object]] = {}
    for record in payload:
        if not isinstance(record, dict):
            raise ValueError("Every reference must be a JSON object.")
        reference_id = record.get("id")
        authors = record.get("authors")
        if (
            not isinstance(reference_id, str)
            or not reference_id
            or not isinstance(authors, list)
            or not authors
            or any(not isinstance(author, str) or not author for author in authors)
            or not isinstance(record.get("year"), int)
        ):
            raise ValueError("Reference metadata is incomplete or malformed.")
        if reference_id in references:
            raise ValueError(f"Duplicate reference ID: {reference_id}")
        references[reference_id] = record
    return references


def load_summary(path: Path) -> dict[str, dict[str, str]]:
    """Load summary.csv as one uniquely keyed row per scenario."""

    with path.open(newline="", encoding="utf-8") as stream:
        reader = csv.DictReader(stream)
        if not reader.fieldnames or "scenario" not in reader.fieldnames:
            raise ValueError("summary.csv must contain a scenario column.")
        rows = list(reader)
    summary: dict[str, dict[str, str]] = {}
    for row in rows:
        scenario = row.get("scenario", "")
        if not scenario or scenario in summary:
            raise ValueError(f"Summary scenarios must be non-empty and unique: {scenario!r}")
        summary[scenario] = row
    if not summary:
        raise ValueError("summary.csv has no data rows.")
    return summary


def resolve_value_marker(
    marker: str, summary: dict[str, dict[str, str]]
) -> str:
    """Resolve one complete VALUE marker using Python's numeric format mini-language."""

    match = VALUE_PATTERN.fullmatch(marker)
    if not match:
        raise ValueError(f"Malformed VALUE marker: {marker}")
    if marker not in CANONICAL_VALUE_MARKERS:
        raise ValueError(f"Noncanonical VALUE marker: {marker}")
    scenario, column, output_format = match.groups()
    if scenario not in summary:
        raise ValueError(f"Unknown VALUE marker scenario: {scenario}")
    row = summary[scenario]
    if column not in row or column in {"scenario", "group"}:
        raise ValueError(f"Unknown numeric VALUE marker column: {column}")
    try:
        value = float(row[column])
        if not math.isfinite(value):
            raise ValueError("VALUE marker resolved to a non-finite number.")
        return format(value, output_format)
    except (TypeError, ValueError) as error:
        raise ValueError(f"Invalid VALUE marker: {marker}") from error


def parse_figure_marker(marker: str) -> tuple[str, str, str]:
    """Parse and validate one complete figure marker."""

    match = FIGURE_PATTERN.fullmatch(marker)
    if not match:
        raise ValueError(f"Malformed FIGURE marker: {marker}")
    if marker not in CANONICAL_FIGURE_MARKERS:
        raise ValueError(f"Noncanonical FIGURE marker: {marker}")
    filename, caption, bookmark = (value.strip() for value in match.groups())
    if (
        Path(filename).name != filename
        or Path(filename).suffix.lower() != ".png"
        or not caption
        or not re.fullmatch(r"fig[1-5]", bookmark)
    ):
        raise ValueError(f"Unsafe or malformed FIGURE marker: {marker}")
    return filename, caption, bookmark


def _author_surname(author: str) -> str:
    return author.strip().split()[-1]


def resolve_citation_marker(
    marker: str, references: dict[str, dict[str, object]]
) -> str:
    """Render one citation marker as a parenthetical Harvard author-date citation."""

    match = CITATION_PATTERN.fullmatch(marker)
    if not match:
        raise ValueError(f"Malformed citation marker: {marker}")
    reference_id = match.group(1)
    if reference_id not in references:
        raise ValueError(f"Unknown citation key: {reference_id}")
    record = references[reference_id]
    authors = record["authors"]
    assert isinstance(authors, list)
    surnames = [_author_surname(str(author)) for author in authors]
    if len(surnames) == 1:
        author_text = surnames[0]
    elif len(surnames) == 2:
        author_text = f"{surnames[0]} and {surnames[1]}"
    else:
        author_text = f"{surnames[0]} et al."
    return f"({author_text}, {record['year']})"


def _normalize_inline_math(text: str) -> str:
    def normalize_span(match: re.Match[str]) -> str:
        expression = match.group(1)
        for source, target in {
            r"\alpha": "alpha",
            r"\beta": "beta",
            r"\gamma": "gamma",
            r"\delta": "delta",
            r"\tau": "tau",
            "{,}": ",",
            "{:}": ":",
        }.items():
            expression = expression.replace(source, target)
        expression = re.sub(r"_\{([A-Za-z0-9]+)\}", r"_\1", expression)
        if "\\" in expression or "{" in expression or "}" in expression:
            raise ValueError(f"Unsupported inline math: {match.group(0)}")
        return expression

    normalized = re.sub(r"\\\((.*?)\\\)", normalize_span, text)
    if r"\(" in normalized or r"\)" in normalized:
        raise ValueError(f"Unsupported inline math: {text}")
    return normalized


def resolve_inline_markup(
    paragraph: str, references: dict[str, dict[str, object]]
) -> list[tuple[str, bool, bool]]:
    """Resolve citations and the project's constrained bold/italic inline markup."""

    if "\n" in paragraph or "\r" in paragraph:
        raise ValueError("Inline-markup resolver accepts one paragraph at a time.")
    paragraph = CITATION_PATTERN.sub(
        lambda match: resolve_citation_marker(match.group(0), references), paragraph
    )
    if "[@" in paragraph:
        raise ValueError(f"Malformed citation marker remains: {paragraph}")
    paragraph = _normalize_inline_math(paragraph)
    if paragraph.count("`") % 2:
        raise ValueError(f"Unmatched inline-code delimiter: {paragraph}")
    pieces: list[tuple[str, bool, bool]] = []
    cursor = 0
    markup = re.compile(
        r"\*\*([^*\r\n]+)\*\*|\*([^*\r\n]+)\*|`([^`\r\n]+)`"
    )
    for match in markup.finditer(paragraph):
        if match.start() > cursor:
            pieces.append((paragraph[cursor : match.start()], False, False))
        if match.group(1) is not None:
            pieces.append((match.group(1), True, False))
        elif match.group(2) is not None:
            pieces.append((match.group(2), False, True))
        else:
            pieces.append((match.group(3), False, False))
        cursor = match.end()
    if cursor < len(paragraph):
        pieces.append((paragraph[cursor:], False, False))
    if any("*" in text or "`" in text for text, _, _ in pieces):
        raise ValueError(f"Malformed or unsupported inline emphasis: {paragraph}")
    return [piece for piece in pieces if piece[0]]


def validate_source_contract(
    text: str,
    references: dict[str, dict[str, object]],
    summary: dict[str, dict[str, str]],
) -> None:
    """Reject anything outside the validated, canonical Markdown contract."""

    for line_number, raw_line in enumerate(text.splitlines(), start=1):
        if "\t" in raw_line:
            raise ValueError(f"Tabs are unsupported on source line {line_number}.")
        if raw_line and raw_line != raw_line.lstrip():
            raise ValueError(
                f"Indented or nested Markdown is unsupported on source line {line_number}."
            )
        line = raw_line.strip()
        if not line:
            continue
        if line.startswith((">", "```", "~~~")):
            raise ValueError(f"Unsupported Markdown on source line {line_number}: {line}")
        if re.match(r"^\d+[.)]\s+", line) or re.fullmatch(r"[-*_]{3,}", line):
            raise ValueError(f"Unsupported Markdown on source line {line_number}: {line}")
        if re.search(r"!?\[[^\]\r\n]+\]\([^)]+\)|~~", line):
            raise ValueError(f"Unsupported Markdown on source line {line_number}: {line}")
        if line.startswith("#") and not re.fullmatch(r"#{1,4}\s+\S.*", line):
            raise ValueError(f"Malformed heading on source line {line_number}: {line}")
        if line.startswith("-") and not re.fullmatch(r"-\s+\S.*", line):
            raise ValueError(f"Malformed bullet on source line {line_number}: {line}")
        if line.count("`") % 2:
            raise ValueError(
                f"Unmatched inline-code delimiter on source line {line_number}."
            )

    scrubbed = DOUBLE_BRACKET_PATTERN.sub("", text)
    if "[[" in scrubbed or "]]" in scrubbed:
        raise ValueError("Unmatched double-bracket marker delimiter in source.")
    source_values = Counter(
        match.group(0) for match in VALUE_PATTERN.finditer(text)
    )
    canonical_values = Counter(
        match.group(0) for match in VALUE_PATTERN.finditer(CANONICAL_SOURCE_TEXT)
    )
    if source_values != canonical_values:
        raise ValueError("Source VALUE markers do not match the canonical manuscript.")
    source_figures = Counter(
        match.group(0) for match in FIGURE_PATTERN.finditer(text)
    )
    canonical_figures = Counter(
        match.group(0) for match in FIGURE_PATTERN.finditer(CANONICAL_SOURCE_TEXT)
    )
    if source_figures != canonical_figures:
        raise ValueError("Source FIGURE markers do not match the canonical manuscript.")

    try:
        from scripts import validate_report_draft as draft_validator
    except ImportError:
        import validate_report_draft as draft_validator

    reference_ids = set(references)
    summary_rows = list(summary.values())
    summary_columns = list(next(iter(summary_rows)).keys())
    positions = draft_validator.heading_positions(text)
    draft_validator.validate_figures(text)
    draft_validator.validate_fixed_markers(text)
    draft_validator.validate_citations_and_references(text, reference_ids)
    draft_validator.validate_values(text, summary_columns, summary_rows)
    draft_validator.validate_equations(text)
    draft_validator.validate_marker_grammar(text)
    draft_validator.validate_language(text)
    word_count = draft_validator.count_main_text_words(text, positions)
    if not 4_500 <= word_count <= 5_500:
        raise ValueError(
            f"Canonical main-text word count must be 4,500-5,500; found {word_count}."
        )


def _author_reference_text(authors: list[str]) -> str:
    formatted: list[str] = []
    for author in authors:
        words = author.split()
        surname = words[-1]
        initials = " ".join(f"{word[0]}." for word in words[:-1] if word)
        formatted.append(f"{surname}, {initials}")
    if len(formatted) == 1:
        return formatted[0]
    return ", ".join(formatted[:-1]) + f" and {formatted[-1]}"


def format_reference(record: dict[str, object]) -> str:
    """Format one canonical record as a compact Harvard-style reference."""

    authors = record.get("authors")
    if not isinstance(authors, list) or not authors:
        raise ValueError("Reference has no ordered author list.")
    author_text = _author_reference_text([str(author) for author in authors])
    year = record.get("year")
    title = str(record.get("title", "")).strip()
    journal = str(record.get("journal", "")).strip()
    volume = str(record.get("volume", "")).strip()
    issue = str(record.get("issue", "")).strip()
    pages = str(record.get("pages_or_article", "")).strip()
    url = str(record.get("url", "")).strip()
    if not title or not journal or not volume or not pages or not url:
        raise ValueError("Reference metadata is incomplete.")
    volume_issue = volume + (f"({issue})" if issue else "")
    return (
        f"{author_text} ({year}) '{title}', {journal}, "
        f"{volume_issue}, {pages}. {url}"
    )


def _set_run_font(
    run,
    *,
    name: str,
    size_pt: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    if size_pt is not None:
        run.font.size = Pt(size_pt)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def _configure_style_font(
    style,
    *,
    name: str,
    size_pt: float,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    style.font.name = name
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), name)
    style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), name)
    style.font.size = Pt(size_pt)
    if color is not None:
        style.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        style.font.bold = bold
    if italic is not None:
        style.font.italic = italic


def _remove_paragraph_borders(paragraph_properties) -> None:
    for borders in list(paragraph_properties.findall(qn("w:pBdr"))):
        paragraph_properties.remove(borders)


def _configure_styles(document: DocumentObject) -> None:
    styles = document.styles
    normal = styles["Normal"]
    _configure_style_font(
        normal, name=PRESET["body_font"], size_pt=PRESET["body_size_pt"]
    )
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(PRESET["body_after_pt"])
    normal.paragraph_format.line_spacing = PRESET["body_line_spacing"]

    for style_name, token in (
        ("Heading 1", PRESET["h1"]),
        ("Heading 2", PRESET["h2"]),
        ("Heading 3", PRESET["h3"]),
    ):
        size, color, before, after = token
        style = styles[style_name]
        _configure_style_font(
            style,
            name=PRESET["body_font"],
            size_pt=size,
            color=color,
            bold=True,
        )
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    caption = styles["Caption"]
    _configure_style_font(
        caption,
        name=CAPTION_STYLE["font"],
        size_pt=CAPTION_STYLE["size_pt"],
        color=CAPTION_STYLE["color"],
        italic=True,
    )
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_before = Pt(CAPTION_STYLE["before_pt"])
    caption.paragraph_format.space_after = Pt(CAPTION_STYLE["after_pt"])
    caption.paragraph_format.line_spacing = CAPTION_STYLE["line_spacing"]
    caption.paragraph_format.keep_together = True

    title = styles["Title"]
    _configure_style_font(
        title,
        name=PRESET["body_font"],
        size_pt=EDITORIAL_COVER["title_size_pt"],
        color=EDITORIAL_COVER["title_color"],
        bold=True,
    )
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    _remove_paragraph_borders(title._element.get_or_add_pPr())

    subtitle = styles["Subtitle"]
    _configure_style_font(
        subtitle,
        name=PRESET["body_font"],
        size_pt=EDITORIAL_COVER["subtitle_size_pt"],
        color=EDITORIAL_COVER["subtitle_color"],
        italic=False,
    )
    subtitle.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(32)

    if "Equation" not in styles:
        equation = styles.add_style("Equation", WD_STYLE_TYPE.PARAGRAPH)
        equation.base_style = normal
    equation = styles["Equation"]
    _configure_style_font(equation, name="Cambria Math", size_pt=11)
    equation.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    equation.paragraph_format.space_before = Pt(4)
    equation.paragraph_format.space_after = Pt(8)
    equation.paragraph_format.line_spacing = 1.15
    equation.paragraph_format.keep_together = True

    if "Cover Metadata" not in styles:
        cover_metadata = styles.add_style("Cover Metadata", WD_STYLE_TYPE.PARAGRAPH)
        cover_metadata.base_style = normal
    cover_metadata = styles["Cover Metadata"]
    _configure_style_font(cover_metadata, name=PRESET["body_font"], size_pt=11)
    cover_metadata.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover_metadata.paragraph_format.space_before = Pt(0)
    cover_metadata.paragraph_format.space_after = Pt(5)
    cover_metadata.paragraph_format.line_spacing = 1.15

    if "Reproducibility Command" not in styles:
        command = styles.add_style(
            "Reproducibility Command", WD_STYLE_TYPE.PARAGRAPH
        )
        command.base_style = normal
    command = styles["Reproducibility Command"]
    _configure_style_font(command, name="Consolas", size_pt=9.5, color="1F4D78")
    command.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    command.paragraph_format.left_indent = Inches(0.25)
    command.paragraph_format.right_indent = Inches(0.25)
    command.paragraph_format.space_before = Pt(4)
    command.paragraph_format.space_after = Pt(8)
    command.paragraph_format.line_spacing = 1.0


def _next_numbering_id(document: DocumentObject) -> tuple[int, int]:
    numbering = document.part.numbering_part.element
    abstract_ids = [
        int(element.get(qn("w:abstractNumId")))
        for element in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(element.get(qn("w:numId")))
        for element in numbering.findall(qn("w:num"))
    ]
    return max(abstract_ids, default=0) + 1, max(num_ids, default=0) + 1


def _configure_real_bullet_style(document: DocumentObject) -> str:
    styles = document.styles
    if "Narrative Bullet" not in styles:
        bullet_style = styles.add_style("Narrative Bullet", WD_STYLE_TYPE.PARAGRAPH)
        bullet_style.base_style = styles["Normal"]
    bullet_style = styles["Narrative Bullet"]
    _configure_style_font(
        bullet_style, name=PRESET["body_font"], size_pt=PRESET["body_size_pt"]
    )
    bullet_style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    bullet_style.paragraph_format.space_before = Pt(0)
    bullet_style.paragraph_format.space_after = Pt(4)
    bullet_style.paragraph_format.line_spacing = 1.208
    bullet_style.paragraph_format.left_indent = Twips(540)
    bullet_style.paragraph_format.first_line_indent = Twips(-280)

    abstract_id, num_id = _next_numbering_id(document)
    numbering = document.part.numbering_part.element
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    num_format = OxmlElement("w:numFmt")
    num_format.set(qn("w:val"), "bullet")
    level.append(num_format)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), "\u2022")
    level.append(level_text)
    level_justification = OxmlElement("w:lvlJc")
    level_justification.set(qn("w:val"), "left")
    level.append(level_justification)
    level_ppr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), "540")
    tabs.append(tab)
    level_ppr.append(tabs)
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "540")
    indent.set(qn("w:hanging"), "280")
    level_ppr.append(indent)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "80")
    spacing.set(qn("w:line"), "290")
    spacing.set(qn("w:lineRule"), "auto")
    level_ppr.append(spacing)
    level.append(level_ppr)
    level_rpr = OxmlElement("w:rPr")
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), PRESET["body_font"])
    fonts.set(qn("w:hAnsi"), PRESET["body_font"])
    level_rpr.append(fonts)
    level.append(level_rpr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    num.append(abstract_reference)
    numbering.append(num)

    ppr = bullet_style._element.get_or_add_pPr()
    existing_num_pr = ppr.find(qn("w:numPr"))
    if existing_num_pr is not None:
        ppr.remove(existing_num_pr)
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_pr.append(ilvl)
    num_id_element = OxmlElement("w:numId")
    num_id_element.set(qn("w:val"), str(num_id))
    num_pr.append(num_id_element)
    ppr.insert(0, num_pr)
    return bullet_style.name


def _configure_page(document: DocumentObject) -> None:
    document.settings.odd_and_even_pages_header_footer = True
    section = document.sections[0]
    section.page_width = Inches(PRESET["page_width_inches"])
    section.page_height = Inches(PRESET["page_height_inches"])
    section.top_margin = Inches(PRESET["margin_inches"])
    section.right_margin = Inches(PRESET["margin_inches"])
    section.bottom_margin = Inches(PRESET["margin_inches"])
    section.left_margin = Inches(PRESET["margin_inches"])
    section.header_distance = Inches(PRESET["header_footer_inches"])
    section.footer_distance = Inches(PRESET["header_footer_inches"])
    section.different_first_page_header_footer = True

    header = section.header
    header.is_linked_to_previous = False
    header_paragraph = header.paragraphs[0]
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_paragraph.paragraph_format.space_after = Pt(0)
    header_run = header_paragraph.add_run(RUNNING_HEADER)
    _set_run_font(
        header_run, name=PRESET["body_font"], size_pt=9, color="6B7280"
    )

    even_header = section.even_page_header
    even_header.is_linked_to_previous = False
    even_header_paragraph = even_header.paragraphs[0]
    even_header_paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    even_header_paragraph.paragraph_format.space_after = Pt(0)
    even_header_run = even_header_paragraph.add_run(RUNNING_HEADER)
    _set_run_font(
        even_header_run, name=PRESET["body_font"], size_pt=9, color="6B7280"
    )

    first_header = section.first_page_header
    first_header.is_linked_to_previous = False
    first_header.paragraphs[0].text = ""

    footer = section.footer
    footer.is_linked_to_previous = False
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_paragraph.paragraph_format.space_after = Pt(0)
    _add_field(footer_paragraph, "PAGE", "2", font_name=PRESET["body_font"])

    even_footer = section.even_page_footer
    even_footer.is_linked_to_previous = False
    even_footer_paragraph = even_footer.paragraphs[0]
    even_footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    even_footer_paragraph.paragraph_format.space_after = Pt(0)
    _add_field(even_footer_paragraph, "PAGE", "2", font_name=PRESET["body_font"])

    first_footer = section.first_page_footer
    first_footer.is_linked_to_previous = False
    first_footer.paragraphs[0].text = ""


def _add_field(
    paragraph, instruction: str, cached_text: str, *, font_name: str
) -> list:
    runs = []
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)
    runs.append(begin_run)

    instruction_run = paragraph.add_run()
    instruction_text = OxmlElement("w:instrText")
    instruction_text.set(XML_SPACE, "preserve")
    instruction_text.text = f" {instruction} "
    instruction_run._r.append(instruction_text)
    runs.append(instruction_run)

    separate_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)
    runs.append(separate_run)

    result_run = paragraph.add_run(cached_text)
    runs.append(result_run)

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)
    runs.append(end_run)
    for run in runs:
        _set_run_font(run, name=font_name)
    return runs


def _add_editorial_cover(
    document: DocumentObject, title: str, metadata: list[tuple[str, str]]
) -> None:
    spacer = document.add_paragraph()
    spacer.paragraph_format.space_before = Pt(EDITORIAL_COVER["top_whitespace_pt"])
    spacer.paragraph_format.space_after = Pt(0)
    spacer.add_run("")

    title_paragraph = document.add_paragraph(title, style="Title")
    title_paragraph.paragraph_format.keep_together = True
    _remove_paragraph_borders(title_paragraph._p.get_or_add_pPr())
    subtitle = document.add_paragraph(
        "Operational Research Case Study", style="Subtitle"
    )
    subtitle.paragraph_format.keep_with_next = True

    for label, value in metadata:
        paragraph = document.add_paragraph(style="Cover Metadata")
        label_run = paragraph.add_run(f"{label}: ")
        _set_run_font(
            label_run,
            name=PRESET["body_font"],
            size_pt=11,
            color="374151",
            bold=True,
        )
        value_run = paragraph.add_run(value)
        _set_run_font(
            value_run,
            name=PRESET["body_font"],
            size_pt=11,
            color="374151",
        )
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _add_rich_paragraph(
    document: DocumentObject,
    text: str,
    references: dict[str, dict[str, object]],
    *,
    style: str | None = None,
):
    paragraph = document.add_paragraph(style=style)
    for segment, bold, italic in resolve_inline_markup(text, references):
        run = paragraph.add_run(segment)
        run.bold = bold
        run.italic = italic
    return paragraph


def _resolve_values(
    text: str, summary: dict[str, dict[str, str]]
) -> str:
    resolved = VALUE_PATTERN.sub(
        lambda match: resolve_value_marker(match.group(0), summary), text
    )
    if "[[VALUE:" in resolved:
        raise ValueError(f"Malformed VALUE marker remains: {text}")
    return resolved


def _add_equation(document: DocumentObject, equation_id: str) -> None:
    if equation_id not in EQUATIONS:
        raise ValueError(f"Unknown equation ID: {equation_id}")
    paragraph = document.add_paragraph(style="Equation")
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for index, line in enumerate(EQUATIONS[equation_id]):
        run = paragraph.add_run(line)
        _set_run_font(run, name="Cambria Math", size_pt=11)
        if index < len(EQUATIONS[equation_id]) - 1:
            run.add_break()


def _add_caption(
    document: DocumentObject, caption: str, bookmark: str, index: int
) -> None:
    paragraph = document.add_paragraph(style="Caption")
    paragraph.add_run("Figure ")
    bookmark_id = str(100 + index)
    bookmark_start = OxmlElement("w:bookmarkStart")
    bookmark_start.set(qn("w:id"), bookmark_id)
    bookmark_start.set(qn("w:name"), bookmark)
    paragraph._p.append(bookmark_start)
    _add_field(paragraph, "SEQ Figure", "0", font_name=CAPTION_STYLE["font"])
    bookmark_end = OxmlElement("w:bookmarkEnd")
    bookmark_end.set(qn("w:id"), bookmark_id)
    paragraph._p.append(bookmark_end)
    paragraph.add_run(f". {caption}")


def _picture_dimensions(path: Path) -> tuple[float, float]:
    with Image.open(path) as image:
        width_pixels, height_pixels = image.size
    aspect = width_pixels / height_pixels
    max_width = 6.35
    max_height = 4.65
    width = max_width
    height = width / aspect
    if height > max_height:
        height = max_height
        width = height * aspect
    return width, height


def _add_figure(
    document: DocumentObject,
    path: Path,
    caption: str,
    bookmark: str,
    index: int,
) -> None:
    if not path.is_file() or path.stat().st_size == 0:
        raise ValueError(f"Missing or empty figure: {path}")
    width, height = _picture_dimensions(path)
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.keep_together = True
    run = paragraph.add_run()
    inline_shape = run.add_picture(
        str(path), width=Inches(width), height=Inches(height)
    )
    doc_properties = inline_shape._inline.docPr
    doc_properties.set("name", f"Figure {index}")
    doc_properties.set("descr", caption)
    _add_caption(document, caption, bookmark, index)


def _set_repeat_table_header(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    row_properties.append(repeat)


def _set_cell_fill(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)
    shading.set(qn("w:val"), "clear")


def _ensure_child(parent, tag: str):
    child = parent.find(qn(tag))
    if child is None:
        child = OxmlElement(tag)
        parent.append(child)
    return child


def _set_width(parent, tag: str, width_dxa: int) -> None:
    width = _ensure_child(parent, tag)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(int(width_dxa)))


def _apply_table_geometry(table, widths: Sequence[int]) -> None:
    if len(widths) != len(table.columns) or sum(widths) != PRESET["content_width_dxa"]:
        raise ValueError("Table widths must match all columns and sum to 9360 DXA.")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table_properties = table._tbl.tblPr
    _set_width(table_properties, "w:tblW", PRESET["content_width_dxa"])
    indent = _ensure_child(table_properties, "w:tblInd")
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(PRESET["table_indent_dxa"]))
    layout = _ensure_child(table_properties, "w:tblLayout")
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        grid_column = OxmlElement("w:gridCol")
        grid_column.set(qn("w:w"), str(width))
        grid.append(grid_column)

    top, bottom, start, end = PRESET["table_cell_margins_dxa"]
    margins = {"top": top, "bottom": bottom, "start": start, "end": end}
    for column_index, width in enumerate(widths):
        table.columns[column_index].width = Twips(width)
    for row in table.rows:
        row.height = None
        for column_index, cell in enumerate(row.cells):
            width = widths[column_index]
            cell.width = Twips(width)
            properties = cell._tc.get_or_add_tcPr()
            _set_width(properties, "w:tcW", width)
            cell_margins = _ensure_child(properties, "w:tcMar")
            for side, margin_width in margins.items():
                margin = _ensure_child(cell_margins, f"w:{side}")
                margin.set(qn("w:w"), str(margin_width))
                margin.set(qn("w:type"), "dxa")


def _format_table(document: DocumentObject, widths: Sequence[int]) -> None:
    table = document.tables[-1]
    table.style = "Table Grid"
    _apply_table_geometry(table, widths)
    _set_repeat_table_header(table.rows[0])
    for row_index, row in enumerate(table.rows):
        for column_index, cell in enumerate(row.cells):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                _set_cell_fill(cell, PRESET["table_header_fill"])
            for paragraph in cell.paragraphs:
                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.LEFT
                    if column_index in (0, len(row.cells) - 1)
                    else WD_ALIGN_PARAGRAPH.CENTER
                )
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                paragraph.paragraph_format.keep_together = True
                for run in paragraph.runs:
                    _set_run_font(
                        run,
                        name=PRESET["body_font"],
                        size_pt=9,
                        bold=row_index == 0,
                    )
    following = document.add_paragraph()
    following.paragraph_format.space_before = Pt(0)
    following.paragraph_format.space_after = Pt(0)


def _read_baseline_parameters() -> dict[str, str]:
    with PARAMETERS_PATH.open(newline="", encoding="utf-8") as stream:
        rows = list(csv.DictReader(stream))
    matches = [row for row in rows if row.get("scenario") == "behaviour_alpha_1"]
    if len(matches) != 1:
        raise ValueError("parameters.csv must contain one behaviour_alpha_1 row.")
    return matches[0]


def _compact_number(value: str, decimals: int = 2) -> str:
    number = float(value)
    if number.is_integer():
        return f"{int(number):,}"
    return f"{number:.{decimals}f}"


def _add_parameter_table(document: DocumentObject) -> None:
    row = _read_baseline_parameters()
    gamma_p = float(row["gamma_p"])
    gamma_w = float(row["gamma_w"])
    delta_p = float(row["delta_p"])
    initial = (
        f"({float(row['initial_susceptible']):.2f}, "
        f"{float(row['initial_infected']):.2f}, "
        f"{float(row['initial_worried']):.2f}, "
        f"{float(row['initial_recovered']):.2f})"
    )
    rows = [
        ("Behavioural multiplier, alpha", _compact_number(row["alpha"], 1), "Literature-derived scenario"),
        ("Infection contact rate, beta_P", _compact_number(row["beta_p"]), "Literature-derived model value"),
        ("Worried-well contact rate, beta_W", _compact_number(row["beta_w"]), "Literature-derived model value"),
        ("Cross-contact rate, beta_WP", _compact_number(row["beta_wp"]), "Literature-derived model value"),
        ("Infected recovery rate, gamma_P", f"1/14 ({gamma_p:.4f})", "Literature-derived model value"),
        ("Worried-well relaxation, gamma_W", f"1/14 ({gamma_w:.4f})", "Literature-derived model value"),
        ("Loss of recovered status, delta_P", f"1/240 ({delta_p:.4f})", "Literature-derived model value"),
        ("Initial state (S, I_P, I_W, R_P)", initial, "Literature-derived baseline"),
        ("Synthetic population, N", f"{int(float(row['population'])):,}", "Coursework assumption"),
        ("Service seeking (p_P, p_W)", f"({_compact_number(row['p_infected'])}, {_compact_number(row['p_worried'])})", "Coursework assumption"),
        ("Coefficient of variation", _compact_number(row["cv"]), "Coursework assumption"),
        ("Mismatch costs (c_u, c_o)", f"({_compact_number(row['underage'])}, {_compact_number(row['overage'])})", "Coursework assumption"),
        (
            "Scenario samples (fit / evaluate)",
            f"{int(float(row['in_sample'])):,} / {int(float(row['out_sample'])):,}",
            f"Seeds {row['in_sample_seed']} / {row['out_sample_seed']}",
        ),
    ]
    table = document.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, ("Parameter", "Baseline value", "Provenance")):
        cell.text = value
    for parameter, value, provenance in rows:
        cells = table.add_row().cells
        for cell, text in zip(cells, (parameter, value, provenance)):
            cell.text = text
    _format_table(document, [3000, 1900, 4460])


def _add_policy_table(
    document: DocumentObject, summary: dict[str, dict[str, str]]
) -> None:
    selected = [
        ("cost_cu_1_co_1", "Balanced 1:1"),
        ("cost_cu_5_co_1", "Baseline 5:1"),
        ("cost_cu_10_co_1", "High-shortage 10:1"),
    ]
    table = document.add_table(rows=1, cols=4)
    headings = (
        "Cost setting",
        "Optimized cost",
        "Mean-policy cost",
        "Signed reduction (%)",
    )
    for cell, value in zip(table.rows[0].cells, headings):
        cell.text = value
    for scenario, label in selected:
        if scenario not in summary:
            raise ValueError(f"Missing policy-table summary scenario: {scenario}")
        row = summary[scenario]
        values = (
            label,
            f"{float(row['optimized_total_cost']):,.1f}",
            f"{float(row['mean_policy_total_cost']):,.1f}",
            f"{float(row['relative_cost_reduction_pct']):.4f}",
        )
        cells = table.add_row().cells
        for cell, value in zip(cells, values):
            cell.text = value
    _format_table(document, [2460, 2100, 2280, 2520])


def _extract_cover(source_lines: list[str]) -> tuple[str, list[tuple[str, str]], int]:
    if not source_lines or not source_lines[0].startswith("# "):
        raise ValueError("Report source must begin with a level-one title.")
    title = source_lines[0][2:].strip()
    metadata: list[tuple[str, str]] = []
    body_start = -1
    metadata_pattern = re.compile(r"^\*\*([^*]+):\*\*\s*(.+)$")
    for index, line in enumerate(source_lines[1:], start=1):
        if line.startswith("## "):
            body_start = index
            break
        if not line.strip():
            continue
        match = metadata_pattern.fullmatch(line)
        if not match:
            raise ValueError(f"Malformed cover metadata line: {line}")
        metadata.append((match.group(1).strip(), match.group(2).strip()))
    if body_start < 0:
        raise ValueError("Report source contains no body sections.")
    if [label for label, _ in metadata] != [
        "Student",
        "Student ID",
        "Module",
        "Facilitator",
        "Submission date",
    ]:
        raise ValueError("Cover metadata labels or order do not match the manuscript contract.")
    return title, metadata, body_start


def _validate_final_marker_state(document: DocumentObject) -> None:
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    text += "\n" + "\n".join(
        cell.text for table in document.tables for row in table.rows for cell in row.cells
    )
    tokens = DOUBLE_BRACKET_PATTERN.findall(text)
    if set(tokens) != PERMITTED_IDENTITY_TOKENS or len(tokens) != 2:
        raise ValueError(
            "Only [[STUDENT_NAME]] and [[STUDENT_ID]] may remain after building."
        )


def build_report(source: Path, output: Path) -> None:
    """Build the final DOCX from the constrained Markdown and generated evidence."""

    references = load_references(REFERENCE_PATH)
    summary = load_summary(SUMMARY_PATH)
    source_text = source.read_text(encoding="utf-8")
    validate_source_contract(source_text, references, summary)
    source_lines = source_text.splitlines()
    title, metadata, body_start = _extract_cover(source_lines)

    document = Document()
    document.core_properties.title = title
    document.core_properties.subject = (
        "MATH6186 operational research case-study draft"
    )
    document.core_properties.author = ""
    document.core_properties.keywords = (
        "worried well; consultation capacity; newsvendor; operational research"
    )
    _configure_styles(document)
    bullet_style = _configure_real_bullet_style(document)
    _configure_page(document)
    _add_editorial_cover(document, title, metadata)

    figure_count = 0
    equation_count = 0
    parameter_table_count = 0
    policy_table_count = 0
    reference_marker_count = 0
    for raw_line in source_lines[body_start:]:
        line = raw_line
        if not line:
            continue
        heading_match = re.fullmatch(r"(#{2,4})\s+(.+)", line)
        if heading_match:
            level = len(heading_match.group(1)) - 1
            heading_text = heading_match.group(2)
            if heading_text in {"References", "Appendix A. Reproducibility"}:
                break_paragraph = document.add_paragraph()
                break_paragraph.paragraph_format.space_after = Pt(0)
                break_paragraph.add_run().add_break(WD_BREAK.PAGE)
            document.add_paragraph(heading_text, style=f"Heading {level}")
            continue

        value_resolved = _resolve_values(line, summary)
        equation_match = EQUATION_PATTERN.fullmatch(value_resolved)
        if equation_match:
            _add_equation(document, equation_match.group(1))
            equation_count += 1
            continue
        if value_resolved == "[[PARAMETER_TABLE]]":
            _add_parameter_table(document)
            parameter_table_count += 1
            continue
        if value_resolved == "[[POLICY_TABLE]]":
            _add_policy_table(document, summary)
            policy_table_count += 1
            continue
        figure_match = FIGURE_PATTERN.fullmatch(value_resolved)
        if figure_match:
            filename, caption, bookmark = parse_figure_marker(value_resolved)
            figure_count += 1
            expected_bookmark = f"fig{figure_count}"
            if bookmark != expected_bookmark:
                raise ValueError(
                    f"Figure bookmarks must be sequential: expected {expected_bookmark}"
                )
            _add_figure(
                document,
                FIGURE_DIR / filename,
                caption,
                bookmark,
                figure_count,
            )
            continue

        bullet_match = re.fullmatch(r"-\s+(.+)", value_resolved)
        if bullet_match:
            item_text = bullet_match.group(1)
            reference_match = REFERENCE_PATTERN.fullmatch(item_text)
            if reference_match:
                reference_id = reference_match.group(1)
                if reference_id not in references:
                    raise ValueError(f"Unknown reference marker: {reference_id}")
                item_text = format_reference(references[reference_id])
                reference_marker_count += 1
            elif "[[REFERENCE:" in item_text:
                raise ValueError(f"Malformed reference marker: {item_text}")
            _add_rich_paragraph(
                document,
                item_text,
                references,
                style=bullet_style,
            )
            continue

        if value_resolved.startswith("`") or value_resolved.endswith("`"):
            if not (
                value_resolved.startswith("`")
                and value_resolved.endswith("`")
                and value_resolved.count("`") == 2
            ):
                raise ValueError(f"Malformed command markup: {value_resolved}")
            command = value_resolved[1:-1]
            document.add_paragraph(command, style="Reproducibility Command")
            continue

        unresolved = [
            token
            for token in DOUBLE_BRACKET_PATTERN.findall(value_resolved)
            if token not in PERMITTED_IDENTITY_TOKENS
        ]
        if unresolved:
            raise ValueError(f"Unresolved marker(s): {', '.join(unresolved)}")
        _add_rich_paragraph(document, value_resolved, references)

    if figure_count != 5:
        raise ValueError(f"Expected five figures, built {figure_count}.")
    if equation_count != 4:
        raise ValueError(f"Expected four equations, built {equation_count}.")
    if parameter_table_count != 1 or policy_table_count != 1:
        raise ValueError("Expected one parameter table and one policy table.")
    if reference_marker_count != len(references):
        raise ValueError(
            f"Expected {len(references)} reference markers, built {reference_marker_count}."
        )
    _validate_final_marker_state(document)

    output.parent.mkdir(parents=True, exist_ok=True)
    temporary_output = output.with_suffix(".tmp.docx")
    document.save(temporary_output)
    temporary_output.replace(output)


def main() -> None:
    build_report(DEFAULT_SOURCE, DEFAULT_OUTPUT)
    print(
        f"Built report DOCX: {DEFAULT_OUTPUT.relative_to(ROOT)} "
        f"({DEFAULT_OUTPUT.stat().st_size:,} bytes)"
    )


if __name__ == "__main__":
    main()
