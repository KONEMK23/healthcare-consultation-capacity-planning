"""Structurally validate the generated MATH6186 Word report."""

from __future__ import annotations

import hashlib
import posixpath
import re
import sys
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree


EXPECTED_SECTION_TITLES = [
    "Abstract",
    "1. Introduction",
    "2. Literature Review",
    "3. Deterministic Worried-Well Model",
    "4. Stochastic Consultation Demand",
    "5. Newsvendor Formulation",
    "6. Experimental Design",
    "7. Results",
    "8. Discussion",
    "9. Conclusion",
    "References",
]
EXPECTED_MEDIA_COUNT = 5
EXPECTED_TABLE_COUNT = 2
EXPECTED_CAPTIONS = 5
EXPECTED_EQUATIONS = [
    (
        "dS/dt = -(beta_P + beta_WP)SI_P - beta_W SI_W + delta_P R_P + gamma_W I_W\n"
        "dI_P/dt = beta_P SI_P + alpha beta_P I_P I_W - gamma_P I_P\n"
        "dI_W/dt = -alpha beta_P I_P I_W + beta_W SI_W + beta_WP SI_P - gamma_W I_W\n"
        "dR_P/dt = gamma_P I_P - delta_P R_P"
    ),
    "m_t = N[p_P I_P(t) + p_W I_W(t)]",
    "min_q (1/n) sum_j [c_u(D_j-q)^+ + c_o(q-D_j)^+]",
    "tau = c_u/(c_u+c_o), q* = F_n^{-1}(tau)",
]
EXPECTED_STYLES = {
    "Normal": {
        "font": "Calibri",
        "size": "22",
        "color": None,
        "before": "0",
        "after": "160",
        "line": "320",
        "line_rule": "auto",
        "alignment": "both",
    },
    "Heading 1": {
        "font": "Calibri",
        "size": "32",
        "color": "2E74B5",
        "before": "360",
        "after": "200",
        "line": None,
        "line_rule": None,
        "alignment": None,
    },
    "Heading 2": {
        "font": "Calibri",
        "size": "26",
        "color": "2E74B5",
        "before": "240",
        "after": "120",
        "line": None,
        "line_rule": None,
        "alignment": None,
    },
    "Heading 3": {
        "font": "Calibri",
        "size": "24",
        "color": "1F4D78",
        "before": "160",
        "after": "80",
        "line": None,
        "line_rule": None,
        "alignment": None,
    },
    "Caption": {
        "font": "Calibri",
        "size": "18",
        "color": "595959",
        "before": "80",
        "after": "120",
        "line": "240",
        "line_rule": "auto",
        "alignment": "center",
    },
}
STYLE_IDS = {
    "Normal": "Normal",
    "Heading 1": "Heading1",
    "Heading 2": "Heading2",
    "Heading 3": "Heading3",
    "Caption": "Caption",
}
ROOT = Path(__file__).resolve().parent.parent
DOCX_PATH = ROOT / "report" / "MATH6186_case_study_draft.docx"
SOURCE_PATH = ROOT / "report" / "MATH6186_case_study_draft.md"
REFERENCE_PATH = ROOT / "report" / "references.json"
FIGURE_DIR = ROOT / "outputs" / "figures"
NS = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
    "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
    "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
    "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
    "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
}
IDENTITY_TOKEN = re.compile(r"\[\[[^\[\]\r\n]+\]\]")
PERMITTED_IDENTITY_TOKENS = {"[[STUDENT_NAME]]", "[[STUDENT_ID]]"}
PROHIBITED_MARKERS = (
    "[[VALUE:",
    "[[FIGURE:",
    "[[EQUATION:",
    "[[REFERENCE:",
    "[[PARAMETER_TABLE]]",
    "[[POLICY_TABLE]]",
)


def fail(message: str) -> None:
    raise ValueError(message)


def require_attribute(
    element: etree._Element | None, attribute: str, expected: str, label: str
) -> None:
    if element is None:
        fail(f"{label} is missing.")
    actual = element.get(f"{{{NS['w']}}}{attribute}")
    if actual != expected:
        fail(f"{label} must be {expected}, found {actual!r}.")


def xml_part(archive: zipfile.ZipFile, name: str) -> etree._Element:
    try:
        payload = archive.read(name)
    except KeyError as error:
        fail(f"DOCX is missing required part: {name}")
        raise AssertionError("unreachable") from error
    return etree.fromstring(payload)


def paragraph_text(paragraph: etree._Element) -> str:
    return "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))


def paragraph_style(paragraph: etree._Element) -> str:
    styles = paragraph.xpath("./w:pPr/w:pStyle/@w:val", namespaces=NS)
    return styles[0] if styles else ""


def validate_complex_field(
    paragraph: etree._Element,
    instruction: str,
    cached_display: str,
    label: str,
) -> tuple[int, int]:
    nodes = list(paragraph.iter())
    positions = {id(node): index for index, node in enumerate(nodes)}
    field_chars = paragraph.xpath(".//w:fldChar", namespaces=NS)
    field_types = [
        node.get(f"{{{NS['w']}}}fldCharType") for node in field_chars
    ]
    if field_types != ["begin", "separate", "end"]:
        fail(f"{label} field chars must be ordered begin/separate/end: {field_types!r}")
    begin_position, separate_position, end_position = [
        positions[id(node)] for node in field_chars
    ]
    instructions = paragraph.xpath(".//w:instrText", namespaces=NS)
    if len(instructions) != 1:
        fail(f"{label} must contain exactly one field instruction.")
    instruction_position = positions[id(instructions[0])]
    if not begin_position < instruction_position < separate_position < end_position:
        fail(f"{label} field instruction/result ordering is malformed.")
    if " ".join((instructions[0].text or "").split()) != instruction:
        fail(f"{label} field instruction must be {instruction!r}.")
    result = "".join(
        (node.text or "")
        for node in paragraph.xpath(".//w:t", namespaces=NS)
        if separate_position < positions[id(node)] < end_position
    )
    if result != cached_display:
        fail(f"{label} cached display must be {cached_display!r}, found {result!r}.")
    return begin_position, end_position


def load_canonical_semantics() -> dict[str, object]:
    try:
        from scripts import build_report_docx as builder
    except ImportError:
        import build_report_docx as builder

    source_text = SOURCE_PATH.read_text(encoding="utf-8")
    lines = source_text.splitlines()
    references = builder.load_references(REFERENCE_PATH)
    summary = builder.load_summary(builder.SUMMARY_PATH)
    title, metadata, body_start = builder._extract_cover(lines)
    ordinary: list[str] = []
    reference_texts: list[str] = []
    captions: list[str] = []
    figure_files: list[str] = []
    for raw_line in lines[body_start:]:
        line = raw_line
        if not line or re.fullmatch(r"#{2,4}\s+.+", line):
            continue
        resolved = builder._resolve_values(line, summary)
        equation = builder.EQUATION_PATTERN.fullmatch(resolved)
        if equation or resolved in {"[[PARAMETER_TABLE]]", "[[POLICY_TABLE]]"}:
            continue
        figure = builder.FIGURE_PATTERN.fullmatch(resolved)
        if figure:
            filename, caption, _bookmark = builder.parse_figure_marker(resolved)
            figure_files.append(filename)
            captions.append(caption)
            continue
        bullet = re.fullmatch(r"-\s+(.+)", resolved)
        if bullet:
            item = bullet.group(1)
            reference = builder.REFERENCE_PATTERN.fullmatch(item)
            if reference:
                item = builder.format_reference(references[reference.group(1)])
                reference_texts.append(item)
            ordinary.append(
                "".join(
                    text
                    for text, _bold, _italic in builder.resolve_inline_markup(
                        item, references
                    )
                )
            )
            continue
        if resolved.startswith("`") and resolved.endswith("`"):
            ordinary.append(resolved[1:-1])
            continue
        ordinary.append(
            "".join(
                text
                for text, _bold, _italic in builder.resolve_inline_markup(
                    resolved, references
                )
            )
        )

    expected_document = Document()
    builder._add_parameter_table(expected_document)
    builder._add_policy_table(expected_document, summary)
    tables = [
        [[cell.text for cell in row.cells] for row in table.rows]
        for table in expected_document.tables
    ]
    return {
        "title": title,
        "subtitle": "Operational Research Case Study",
        "metadata": [f"{label}: {value}" for label, value in metadata],
        "ordinary": ordinary,
        "references": reference_texts,
        "captions": captions,
        "figure_files": figure_files,
        "tables": tables,
    }


def validate_sections(document_xml: etree._Element) -> None:
    headings = [
        "".join(paragraph.xpath(".//w:t/text()", namespaces=NS))
        for paragraph in document_xml.xpath(
            ".//w:body/w:p[w:pPr/w:pStyle[@w:val='Heading1']]",
            namespaces=NS,
        )
    ]
    if headings != EXPECTED_SECTION_TITLES:
        fail(
            "Heading 1 section sequence does not match the required report sections: "
            f"{headings!r}"
        )


def validate_page_geometry(document_xml: etree._Element) -> None:
    section_properties = document_xml.xpath(".//w:sectPr", namespaces=NS)
    if not section_properties:
        fail("DOCX has no section properties.")
    for index, section in enumerate(section_properties, start=1):
        page_size = section.find("w:pgSz", NS)
        require_attribute(page_size, "w", "12240", f"Section {index} page width")
        require_attribute(page_size, "h", "15840", f"Section {index} page height")
        page_margin = section.find("w:pgMar", NS)
        for edge in ("top", "right", "bottom", "left"):
            require_attribute(
                page_margin, edge, "1440", f"Section {index} {edge} margin"
            )
        require_attribute(
            page_margin, "header", "708", f"Section {index} header distance"
        )
        require_attribute(
            page_margin, "footer", "708", f"Section {index} footer distance"
        )


def validate_styles(styles_xml: etree._Element) -> None:
    for style_name, expected in EXPECTED_STYLES.items():
        matches = styles_xml.xpath(
            ".//w:style[@w:styleId=$style_id]",
            namespaces=NS,
            style_id=STYLE_IDS[style_name],
        )
        if len(matches) != 1:
            fail(f"Expected exactly one {style_name!r} style, found {len(matches)}.")
        style = matches[0]
        run_properties = style.find("w:rPr", NS)
        if run_properties is None:
            fail(f"{style_name} style has no explicit run properties.")
        fonts = run_properties.find("w:rFonts", NS)
        require_attribute(fonts, "ascii", expected["font"], f"{style_name} font")
        require_attribute(fonts, "hAnsi", expected["font"], f"{style_name} font")
        require_attribute(run_properties.find("w:sz", NS), "val", expected["size"], f"{style_name} size")
        color = run_properties.find("w:color", NS)
        if expected["color"] is None:
            if color is not None and color.get(f"{{{NS['w']}}}val") not in (None, "000000", "auto"):
                fail(f"{style_name} color must be black/default.")
        else:
            require_attribute(color, "val", expected["color"], f"{style_name} color")
        paragraph_properties = style.find("w:pPr", NS)
        if paragraph_properties is None:
            fail(f"{style_name} style has no explicit paragraph properties.")
        spacing = paragraph_properties.find("w:spacing", NS)
        require_attribute(spacing, "before", expected["before"], f"{style_name} before spacing")
        require_attribute(spacing, "after", expected["after"], f"{style_name} after spacing")
        if expected["line"] is not None:
            require_attribute(spacing, "line", expected["line"], f"{style_name} line spacing")
            require_attribute(
                spacing,
                "lineRule",
                expected["line_rule"],
                f"{style_name} line-spacing rule",
            )
        if expected["alignment"] is not None:
            require_attribute(
                paragraph_properties.find("w:jc", NS),
                "val",
                expected["alignment"],
                f"{style_name} alignment",
            )


def validate_border_free_cover(
    styles_xml: etree._Element, document_xml: etree._Element
) -> None:
    title_styles = styles_xml.xpath(
        ".//w:style[@w:styleId='Title']", namespaces=NS
    )
    if len(title_styles) != 1:
        fail("DOCX must contain exactly one Title style.")
    if title_styles[0].xpath("./w:pPr/w:pBdr", namespaces=NS):
        fail("Editorial-cover Title style must not contain paragraph borders.")
    title_paragraphs = document_xml.xpath(
        ".//w:body/w:p[w:pPr/w:pStyle[@w:val='Title']]", namespaces=NS
    )
    if len(title_paragraphs) != 1:
        fail("Editorial cover must contain exactly one Title paragraph.")
    if title_paragraphs[0].xpath("./w:pPr/w:pBdr", namespaces=NS):
        fail("Editorial-cover title paragraph must not contain direct borders.")


def validate_tables(
    document_xml: etree._Element, expected_tables: list[list[list[str]]]
) -> None:
    tables = document_xml.xpath(".//w:body/w:tbl", namespaces=NS)
    if len(tables) != EXPECTED_TABLE_COUNT:
        fail(f"Expected {EXPECTED_TABLE_COUNT} tables, found {len(tables)}.")
    actual_tables = [
        [
            [paragraph_text(cell) for cell in row.xpath("./w:tc", namespaces=NS)]
            for row in table.xpath("./w:tr", namespaces=NS)
        ]
        for table in tables
    ]
    if actual_tables != expected_tables:
        fail("Table text does not match canonical parameters.csv/summary.csv content.")
    for index, table in enumerate(tables, start=1):
        properties = table.find("w:tblPr", NS)
        if properties is None:
            fail(f"Table {index} has no table properties.")
        width = properties.find("w:tblW", NS)
        require_attribute(width, "type", "dxa", f"Table {index} width type")
        require_attribute(width, "w", "9360", f"Table {index} width")
        require_attribute(
            properties.find("w:tblInd", NS), "w", "120", f"Table {index} indent"
        )
        require_attribute(
            properties.find("w:tblInd", NS),
            "type",
            "dxa",
            f"Table {index} indent type",
        )
        require_attribute(
            properties.find("w:tblLayout", NS),
            "type",
            "fixed",
            f"Table {index} layout",
        )
        grid_widths = [
            int(column.get(f"{{{NS['w']}}}w"))
            for column in table.xpath("./w:tblGrid/w:gridCol", namespaces=NS)
        ]
        if not grid_widths or sum(grid_widths) != 9360:
            fail(f"Table {index} grid widths must sum to 9360 DXA: {grid_widths!r}")
        header_markers = table.xpath("./w:tr[1]/w:trPr/w:tblHeader", namespaces=NS)
        if len(header_markers) != 1 or header_markers[0].get(
            f"{{{NS['w']}}}val"
        ) not in {"true", "1"}:
            fail(f"Table {index} first row must be marked as a repeating header.")
        for row_number, row in enumerate(table.xpath("./w:tr", namespaces=NS), start=1):
            cells = row.xpath("./w:tc", namespaces=NS)
            if len(cells) != len(grid_widths):
                fail(
                    f"Table {index} row {row_number} has {len(cells)} cells "
                    f"for {len(grid_widths)} grid columns."
                )
            cell_widths: list[int] = []
            for cell in cells:
                tc_width = cell.find("w:tcPr/w:tcW", NS)
                require_attribute(
                    tc_width,
                    "type",
                    "dxa",
                    f"Table {index} row {row_number} cell width type",
                )
                cell_widths.append(int(tc_width.get(f"{{{NS['w']}}}w")))
                margins = cell.find("w:tcPr/w:tcMar", NS)
                expected_margins = {
                    "top": "80",
                    "bottom": "80",
                    "start": "120",
                    "end": "120",
                }
                for side, expected_width in expected_margins.items():
                    margin = margins.find(f"w:{side}", NS) if margins is not None else None
                    require_attribute(
                        margin,
                        "type",
                        "dxa",
                        f"Table {index} row {row_number} {side} cell-margin type",
                    )
                    require_attribute(
                        margin,
                        "w",
                        expected_width,
                        f"Table {index} row {row_number} {side} cell margin",
                    )
            if cell_widths != grid_widths:
                fail(
                    f"Table {index} row {row_number} cell widths {cell_widths!r} "
                    f"do not match grid {grid_widths!r}."
                )
            if row.find("w:trPr/w:trHeight", NS) is not None:
                fail(f"Table {index} row {row_number} must not use a fixed height.")


def validate_figures(
    archive: zipfile.ZipFile,
    document_xml: etree._Element,
    expected_figure_files: list[str],
) -> None:
    media = [
        name
        for name in archive.namelist()
        if name.startswith("word/media/") and not name.endswith("/")
    ]
    if len(media) != EXPECTED_MEDIA_COUNT:
        fail(f"Expected {EXPECTED_MEDIA_COUNT} media files, found {len(media)}.")
    inline = document_xml.xpath(".//wp:inline", namespaces=NS)
    anchored = document_xml.xpath(".//wp:anchor", namespaces=NS)
    if len(inline) != EXPECTED_MEDIA_COUNT:
        fail(f"Expected {EXPECTED_MEDIA_COUNT} inline drawings, found {len(inline)}.")
    if anchored:
        fail(f"Anchored drawings are prohibited; found {len(anchored)}.")
    relationships = xml_part(archive, "word/_rels/document.xml.rels")
    relationship_targets = {
        relationship.get("Id"): relationship.get("Target")
        for relationship in relationships.xpath(".//pr:Relationship", namespaces=NS)
    }
    actual_hashes: list[str] = []
    for index, drawing in enumerate(inline, start=1):
        extent = drawing.find("wp:extent", NS)
        if extent is None:
            fail(f"Figure {index} has no drawing extent.")
        width_emu = int(extent.get("cx", "0"))
        if width_emu <= 0 or width_emu > int(6.35 * 914400):
            fail(f"Figure {index} width is outside (0, 6.35] inches.")
        blips = drawing.xpath(".//a:blip/@r:embed", namespaces=NS)
        if len(blips) != 1 or blips[0] not in relationship_targets:
            fail(f"Figure {index} has no unique resolvable image relationship.")
        target = relationship_targets[blips[0]]
        if not target:
            fail(f"Figure {index} image relationship has no target.")
        part_name = posixpath.normpath(posixpath.join("word", target))
        try:
            actual_hashes.append(hashlib.sha256(archive.read(part_name)).hexdigest())
        except KeyError as error:
            fail(f"Figure {index} image part is missing: {part_name}")
            raise AssertionError("unreachable") from error
    expected_hashes = [
        hashlib.sha256((FIGURE_DIR / filename).read_bytes()).hexdigest()
        for filename in expected_figure_files
    ]
    if actual_hashes != expected_hashes:
        fail("Embedded image identities/order do not match the five canonical PNGs.")


def validate_captions_and_bookmarks(
    document_xml: etree._Element, expected_captions: list[str]
) -> None:
    captions = document_xml.xpath(
        ".//w:body/w:p[w:pPr/w:pStyle[@w:val='Caption']]",
        namespaces=NS,
    )
    if len(captions) != EXPECTED_CAPTIONS:
        fail(f"Expected {EXPECTED_CAPTIONS} captions, found {len(captions)}.")
    expected_bookmarks = {f"fig{index}" for index in range(1, 6)}
    all_starts = document_xml.xpath(".//w:bookmarkStart", namespaces=NS)
    all_ends = document_xml.xpath(".//w:bookmarkEnd", namespaces=NS)
    figure_starts = [
        start
        for start in all_starts
        if start.get(f"{{{NS['w']}}}name") in expected_bookmarks
    ]
    if len(figure_starts) != 5:
        fail("Figure bookmarks must contain exactly one start for each fig1-fig5.")
    names = [start.get(f"{{{NS['w']}}}name") for start in figure_starts]
    if set(names) != expected_bookmarks or len(names) != len(set(names)):
        fail("Figure bookmark starts must be unique fig1-fig5.")
    end_ids = [end.get(f"{{{NS['w']}}}id") for end in all_ends]
    for index, caption in enumerate(captions, start=1):
        text = paragraph_text(caption)
        expected_text = f"Figure {index}. {expected_captions[index - 1]}"
        if text != expected_text:
            fail(
                f"Caption {index} text must match its canonical marker: "
                f"{expected_text!r}, found {text!r}."
            )
        starts = caption.xpath(
            f"./w:bookmarkStart[@w:name='fig{index}']", namespaces=NS
        )
        if len(starts) != 1:
            fail(f"Caption {index} must contain exactly one fig{index} start.")
        bookmark_id = starts[0].get(f"{{{NS['w']}}}id")
        if not bookmark_id or end_ids.count(bookmark_id) != 1:
            fail(f"Caption {index} bookmark must have one matching end ID.")
        ends = caption.xpath(
            f"./w:bookmarkEnd[@w:id='{bookmark_id}']", namespaces=NS
        )
        if len(ends) != 1:
            fail(f"Caption {index} bookmark end must be in the same caption.")
        begin_position, end_position = validate_complex_field(
            caption, "SEQ Figure", str(index), f"Caption {index}"
        )
        children = list(caption.iter())
        positions = {id(node): position for position, node in enumerate(children)}
        if not (
            positions[id(starts[0])]
            < begin_position
            < end_position
            < positions[id(ends[0])]
        ):
            fail(f"Caption {index} bookmark must wrap the complete SEQ field.")


def validate_equations(document_xml: etree._Element) -> None:
    equation_paragraphs = document_xml.xpath(
        ".//w:body/w:p[w:pPr/w:pStyle[@w:val='Equation']]",
        namespaces=NS,
    )
    equations: list[str] = []
    break_counts: list[int] = []
    for paragraph in equation_paragraphs:
        displayed: list[str] = []
        breaks = 0
        for run in paragraph.xpath("./w:r", namespaces=NS):
            for node in run.iter():
                if node.tag == f"{{{NS['w']}}}t":
                    displayed.append(node.text or "")
                elif node.tag == f"{{{NS['w']}}}br":
                    displayed.append("\n")
                    breaks += 1
        equations.append("".join(displayed))
        break_counts.append(breaks)
    if equations != EXPECTED_EQUATIONS:
        fail(f"Expected the four implemented equations exactly, found {equations!r}.")
    if break_counts != [3, 0, 0, 0]:
        fail(f"Equation line breaks must be actual w:br elements: {break_counts!r}.")
    for index, paragraph in enumerate(equation_paragraphs, start=1):
        require_attribute(
            paragraph.find("w:pPr/w:jc", NS),
            "val",
            "center",
            f"Equation {index} alignment",
        )
        fonts = paragraph.xpath(".//w:rPr/w:rFonts", namespaces=NS)
        if not fonts:
            fail(f"Equation {index} has no explicit run font.")
        for font in fonts:
            if font.get(f"{{{NS['w']}}}ascii") != "Cambria Math":
                fail(f"Equation {index} must use Cambria Math.")


def validate_markers(document_xml: etree._Element) -> None:
    text = "\n".join(document_xml.xpath(".//w:t/text()", namespaces=NS))
    for marker in PROHIBITED_MARKERS:
        if marker in text:
            fail(f"Unresolved marker remains in DOCX: {marker}")
    tokens = IDENTITY_TOKEN.findall(text)
    if set(tokens) != PERMITTED_IDENTITY_TOKENS or len(tokens) != 2:
        fail(
            "The only remaining double-bracket tokens must be exactly "
            "[[STUDENT_NAME]] and [[STUDENT_ID]]."
        )


def validate_content_preservation(
    document_xml: etree._Element, expected: dict[str, object]
) -> None:
    body_paragraphs = document_xml.xpath(".//w:body/w:p", namespaces=NS)
    titles = [
        paragraph_text(paragraph)
        for paragraph in body_paragraphs
        if paragraph_style(paragraph) == "Title"
    ]
    subtitles = [
        paragraph_text(paragraph)
        for paragraph in body_paragraphs
        if paragraph_style(paragraph) == "Subtitle"
    ]
    metadata = [
        paragraph_text(paragraph)
        for paragraph in body_paragraphs
        if paragraph_style(paragraph) == "CoverMetadata"
    ]
    if titles != [expected["title"]]:
        fail("Cover title does not match the canonical Markdown title.")
    if subtitles != [expected["subtitle"]]:
        fail("Cover subtitle does not match the editorial-cover contract.")
    if metadata != expected["metadata"]:
        fail("Cover metadata does not match the canonical Markdown metadata.")

    excluded_styles = {
        "Title",
        "Subtitle",
        "CoverMetadata",
        "Heading1",
        "Heading2",
        "Heading3",
        "Equation",
        "Caption",
    }
    ordinary = [
        paragraph_text(paragraph)
        for paragraph in body_paragraphs
        if paragraph_style(paragraph) not in excluded_styles
        and paragraph_text(paragraph)
    ]
    if ordinary != expected["ordinary"]:
        for index, (actual, wanted) in enumerate(
            zip(ordinary, expected["ordinary"]), start=1
        ):
            if actual != wanted:
                fail(
                    f"Ordinary body paragraph {index} differs from canonical content: "
                    f"{actual!r} != {wanted!r}"
                )
        fail(
            "Ordinary body paragraph count differs from canonical content: "
            f"{len(ordinary)} != {len(expected['ordinary'])}."
        )

    references_heading = next(
        (
            index
            for index, paragraph in enumerate(body_paragraphs)
            if paragraph_style(paragraph) == "Heading1"
            and paragraph_text(paragraph) == "References"
        ),
        None,
    )
    if references_heading is None:
        fail("References heading is missing.")
    reference_paragraphs = [
        paragraph_text(paragraph)
        for paragraph in body_paragraphs[references_heading + 1 :]
        if paragraph_text(paragraph)
    ]
    if reference_paragraphs != expected["references"]:
        fail("References do not match canonical references.json formatting/order.")


def validate_page_furniture(
    archive: zipfile.ZipFile,
    document_xml: etree._Element,
    settings_xml: etree._Element,
) -> None:
    even_and_odd = settings_xml.xpath(".//w:evenAndOddHeaders", namespaces=NS)
    if len(even_and_odd) != 1:
        fail(
            "Document settings must explicitly enable odd/even page furniture "
            "for cross-renderer consistency."
        )
    section = document_xml.find(".//w:body/w:sectPr", NS)
    if section is None or section.find("w:titlePg", NS) is None:
        fail("The section must use a different first page to suppress cover furniture.")
    relationship_root = xml_part(archive, "word/_rels/document.xml.rels")
    relationships = {
        relationship.get("Id"): relationship
        for relationship in relationship_root.xpath(
            ".//pr:Relationship", namespaces=NS
        )
    }

    def referenced_parts(kind: str) -> dict[str, str]:
        references = section.xpath(f"./w:{kind}Reference", namespaces=NS)
        by_type: dict[str, str] = {}
        for reference in references:
            reference_type = reference.get(f"{{{NS['w']}}}type")
            relationship_id = reference.get(f"{{{NS['r']}}}id")
            if (
                reference_type in by_type
                or relationship_id not in relationships
                or reference_type not in {"default", "even", "first"}
            ):
                fail(
                    f"Section {kind} references must be unique "
                    "default/even/first triples."
                )
            relationship = relationships[relationship_id]
            target = relationship.get("Target")
            relationship_type = relationship.get("Type", "")
            if not target or not relationship_type.endswith(f"/{kind}"):
                fail(f"Section {kind} reference has the wrong relationship type.")
            by_type[reference_type] = posixpath.normpath(
                posixpath.join("word", target)
            )
        if set(by_type) != {"default", "even", "first"}:
            fail(
                f"Section must reference exactly default, even, and first "
                f"{kind} parts."
            )
        return by_type

    header_references = referenced_parts("header")
    footer_references = referenced_parts("footer")
    header_parts = [
        name
        for name in archive.namelist()
        if re.fullmatch(r"word/header\d+\.xml", name)
    ]
    footer_parts = [
        name
        for name in archive.namelist()
        if re.fullmatch(r"word/footer\d+\.xml", name)
    ]
    if len(header_parts) != 3 or len(footer_parts) != 3:
        fail("Expected separate default, even, and first-page header/footer parts.")
    if set(header_parts) != set(header_references.values()):
        fail("Header parts must be exactly those referenced by the section.")
    if set(footer_parts) != set(footer_references.values()):
        fail("Footer parts must be exactly those referenced by the section.")

    default_header = xml_part(archive, header_references["default"])
    even_header = xml_part(archive, header_references["even"])
    first_header = xml_part(archive, header_references["first"])
    expected_header = "MATH6186 | Worried-Well Consultation Capacity"
    if "".join(default_header.xpath(".//w:t/text()", namespaces=NS)) != expected_header:
        fail("Default running header text is incorrect.")
    if "".join(even_header.xpath(".//w:t/text()", namespaces=NS)) != expected_header:
        fail("Even-page running header text is incorrect.")
    if first_header.xpath(".//w:t/text() | .//w:instrText/text()", namespaces=NS):
        fail("First-page header must be blank.")

    default_footer = xml_part(archive, footer_references["default"])
    even_footer = xml_part(archive, footer_references["even"])
    first_footer = xml_part(archive, footer_references["first"])
    for label, footer in (("Default", default_footer), ("Even-page", even_footer)):
        footer_paragraphs = footer.xpath(".//w:p", namespaces=NS)
        if len(footer_paragraphs) != 1:
            fail(f"The {label.lower()} page-number footer must contain one paragraph.")
        require_attribute(
            footer_paragraphs[0].find("w:pPr/w:jc", NS),
            "val",
            "right",
            f"{label} page-number footer alignment",
        )
        validate_complex_field(
            footer_paragraphs[0], "PAGE", "2", f"{label} PAGE footer"
        )
    if first_footer.xpath(".//w:t/text() | .//w:instrText/text()", namespaces=NS):
        fail("First-page footer must be blank.")
    if first_footer.xpath(".//w:fldChar", namespaces=NS):
        fail("First-page footer must not contain fields.")


def main() -> None:
    if not DOCX_PATH.is_file():
        fail(f"Missing DOCX: {DOCX_PATH.relative_to(ROOT)}")
    if DOCX_PATH.stat().st_size == 0:
        fail(f"Empty DOCX: {DOCX_PATH.relative_to(ROOT)}")
    Document(DOCX_PATH)
    expected = load_canonical_semantics()
    with zipfile.ZipFile(DOCX_PATH) as archive:
        document_xml = xml_part(archive, "word/document.xml")
        styles_xml = xml_part(archive, "word/styles.xml")
        settings_xml = xml_part(archive, "word/settings.xml")
        validate_sections(document_xml)
        validate_page_geometry(document_xml)
        validate_styles(styles_xml)
        validate_border_free_cover(styles_xml, document_xml)
        validate_tables(document_xml, expected["tables"])
        validate_figures(archive, document_xml, expected["figure_files"])
        validate_captions_and_bookmarks(document_xml, expected["captions"])
        validate_equations(document_xml)
        validate_markers(document_xml)
        validate_content_preservation(document_xml, expected)
        validate_page_furniture(archive, document_xml, settings_xml)
    print(
        "Report DOCX validation passed: "
        f"{len(EXPECTED_SECTION_TITLES)} sections, {EXPECTED_TABLE_COUNT} tables, "
        f"{EXPECTED_MEDIA_COUNT} inline figures, {EXPECTED_CAPTIONS} captions, "
        f"{len(EXPECTED_EQUATIONS)} equations."
    )


if __name__ == "__main__":
    try:
        main()
    except (OSError, ValueError, zipfile.BadZipFile, etree.XMLSyntaxError) as error:
        print(f"Report DOCX validation failed: {error}", file=sys.stderr)
        raise SystemExit(1) from error
